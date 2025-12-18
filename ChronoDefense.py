# ===================== IMPORTS =====================
import cv2
import os
import pickle
import random
import socket
import geocoder
import pandas as pd
import numpy as np
from datetime import datetime, timedelta
from tkinter import Tk, Button, Label, messagebox, filedialog, simpledialog
from sklearn.preprocessing import StandardScaler
from sklearn.model_selection import train_test_split
from sklearn.ensemble import RandomForestClassifier
from sklearn.metrics import accuracy_score
import openpyxl
from openpyxl.styles import PatternFill

# ===================== SYSTEM INFO =====================
def get_system_info():
    ip = socket.gethostbyname(socket.gethostname())
    g = geocoder.ip(ip)
    region = g.country if g.country else "Unknown"
    return ip, region

# ===================== DATA GENERATION =====================
def generate_data(records=500):
    ip, region = get_system_info()
    attacks = ['DDoS', 'Phishing', 'Malware', 'Ransomware']
    ports = [80, 443, 21, 22]
    protocols = ['TCP', 'UDP', 'HTTP']

    data = []
    for _ in range(records):
        impact = random.randint(1, 10)
        success = random.randint(1, 10)
        severity = impact * success

        data.append([
            severity, ip,
            random.choice(attacks),
            random.choice(ports),
            datetime.now() - timedelta(days=random.randint(1, 365)),
            random.randint(10, 300),
            region, impact, success,
            random.choice(protocols)
        ])

    columns = [
        'SeverityScore', 'IPAddress', 'AttackType', 'Port',
        'Timestamp', 'Duration', 'Region',
        'Impact', 'Success', 'Protocol'
    ]
    return pd.DataFrame(data, columns=columns)

# ===================== ML MODEL =====================
historical_data = generate_data()
historical_data.to_csv("cyber_history.csv", index=False)

data = pd.read_csv("cyber_history.csv")
data = pd.get_dummies(data, columns=['AttackType', 'Region', 'Protocol'])

X = data.drop(columns=['SeverityScore', 'IPAddress', 'Timestamp'])
y = data['SeverityScore']

scaler = StandardScaler()
X = scaler.fit_transform(X)

X_train, X_test, y_train, y_test = train_test_split(X, y, test_size=0.2)
model = RandomForestClassifier()
model.fit(X_train, y_train)

print("Model Accuracy:", accuracy_score(y_test, model.predict(X_test)))

# ===================== REPORT SAVE =====================
def save_report(df, filename):
    df.to_excel(filename, index=False)
    wb = openpyxl.load_workbook(filename)
    ws = wb.active

    green = PatternFill("solid", "00FF00")
    yellow = PatternFill("solid", "FFFF00")
    red = PatternFill("solid", "FF0000")

    for row in ws.iter_rows(min_row=2):
        score = row[0].value
        fill = green if score < 40 else yellow if score < 70 else red
        for cell in row:
            cell.fill = fill

    wb.save(filename)

# ===================== FACE AUTH =====================
username = ""
face_data = {}
recognizer = cv2.face.LBPHFaceRecognizer_create()

def save_face():
    with open("face.pkl", "wb") as f:
        pickle.dump((username, face_data), f)

def load_face():
    global username, face_data
    if os.path.exists("face.pkl"):
        with open("face.pkl", "rb") as f:
            username, face_data = pickle.load(f)
            recognizer.train(face_data["samples"],
                             np.arange(len(face_data["samples"])))

# ===================== REGISTER FACE =====================
def register_face():
    global username

    if not username:
        messagebox.showerror("Error", "Enter username")
        return

    cap = cv2.VideoCapture(0)
    cascade = cv2.CascadeClassifier(
        cv2.data.haarcascades + "haarcascade_frontalface_default.xml"
    )

    samples = []
    count = 0

    while True:
        ret, frame = cap.read()
        gray = cv2.cvtColor(frame, cv2.COLOR_BGR2GRAY)
        faces = cascade.detectMultiScale(gray, 1.1, 4)

        for x, y, w, h in faces:
            face = gray[y:y+h, x:x+w]
            face = cv2.resize(face, (200, 200))
            samples.append(face)
            count += 1
            cv2.rectangle(frame, (x,y), (x+w,y+h), (0,255,0), 2)
            cv2.putText(frame, f"{count}/20", (x, y-10),
                        cv2.FONT_HERSHEY_SIMPLEX, 0.8, (0,255,0), 2)

        cv2.imshow("Register Face", frame)
        if cv2.waitKey(1) == 27 or count >= 20:
            break

    cap.release()
    cv2.destroyAllWindows()

    if len(samples) < 10:
        messagebox.showerror("Error", "Face not detected properly")
        return

    face_data["samples"] = samples
    recognizer.train(samples, np.arange(len(samples)))
    save_face()
    messagebox.showinfo("Success", "Face Registered")

# ===================== AUTHENTICATE =====================
def authenticate():
    cap = cv2.VideoCapture(0)
    cascade = cv2.CascadeClassifier(
        cv2.data.haarcascades + "haarcascade_frontalface_default.xml"
    )

    while True:
        ret, frame = cap.read()
        gray = cv2.cvtColor(frame, cv2.COLOR_BGR2GRAY)
        faces = cascade.detectMultiScale(gray, 1.1, 4)

        for x, y, w, h in faces:
            face = gray[y:y+h, x:x+w]
            face = cv2.resize(face, (200, 200))
            label, conf = recognizer.predict(face)

            if conf < 60:
                cap.release()
                cv2.destroyAllWindows()
                messagebox.showinfo("Access Granted", f"Welcome {username}")
                generate_report()
                return
            else:
                messagebox.showerror("Denied", "Authentication Failed")
                cap.release()
                cv2.destroyAllWindows()
                return

        cv2.imshow("Authenticate", frame)
        if cv2.waitKey(1) == 27:
            break

    cap.release()
    cv2.destroyAllWindows()

# ===================== REPORT GENERATION =====================
def generate_report():
    file = filedialog.asksaveasfilename(
        defaultextension=".xlsx",
        filetypes=[("Excel Files", "*.xlsx")]
    )
    if file:
        save_report(historical_data, file)
        messagebox.showinfo("Report Generated",
                            "Cyber Security Threat Report Saved Successfully")

# ===================== GUI =====================
def set_username(root):
    global username
    username = simpledialog.askstring("Username", "Enter username", parent=root)

def start_gui():
    root = Tk()
    root.title("AI-Driven Cyber Security Time Machine")

    Label(root, text="Secure Cyber Threat System").pack(pady=5)
    Button(root, text="Enter Username",
           command=lambda: set_username(root)).pack(pady=5)
    Button(root, text="Register Face",
           command=register_face).pack(pady=5)
    Button(root, text="Authenticate & Generate Report",
           command=authenticate).pack(pady=5)

    root.mainloop()

load_face()
start_gui()
